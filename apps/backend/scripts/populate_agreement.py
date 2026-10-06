"""Populate the supplied Safarcars agreement template for one booking."""

from __future__ import annotations

import argparse
import json
from datetime import datetime
from pathlib import Path

from docx import Document


def fmt_date(value: str | None) -> str:
    if not value:
        return "Not recorded"
    try:
        dt = datetime.fromisoformat(value.replace("Z", "+00:00"))
        return dt.strftime("%d %b %Y, %I:%M %p")
    except ValueError:
        return value


def money(value: object) -> str:
    try:
        return f"₹{float(value):,.0f}"
    except (TypeError, ValueError):
        return "₹0"


def name(person: dict | None) -> str:
    if not person:
        return "Not recorded"
    return " ".join(v for v in [person.get("firstName"), person.get("lastName")] if v).strip() or "Not recorded"


def set_paragraph_text(paragraph, value: str) -> None:
    runs = paragraph.runs
    if runs:
        runs[0].text = value
        for run in runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def replace_paragraph(paragraph, replacements: dict[str, str]) -> None:
    current = "".join(run.text or "" for run in paragraph.runs)
    if not current:
        return
    updated = current
    for old, new in replacements.items():
        updated = updated.replace(old, new)
    if updated != current:
        set_paragraph_text(paragraph, updated)


def cell_text(cell) -> str:
    return " ".join(cell.text.split())


def set_cell(cell, value: str) -> None:
    if cell.paragraphs:
        set_paragraph_text(cell.paragraphs[0], value)
        for paragraph in cell.paragraphs[1:]:
            set_paragraph_text(paragraph, "")
    else:
        cell.text = value


def walk_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", required=True)
    parser.add_argument("--data", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    data = json.loads(Path(args.data).read_text(encoding="utf-8"))
    booking = data["booking"]
    customer = data.get("customer") or {}
    host = data.get("host") or {}
    vehicle = data.get("vehicle") or {}
    trip = data.get("trip") or {}
    policy = data.get("policy") or {}
    business = data.get("business") or {}
    effective = data.get("generatedAt") or booking.get("startDate")
    customer_name = name(customer)
    host_name = name(host)
    vehicle_name = " ".join(v for v in [vehicle.get("make"), vehicle.get("model"), str(vehicle.get("year") or "")] if v).strip()

    replacements = {
        "[Host Full Name]": host_name,
        "[Host name]": host_name,
        "[Guest name]": customer_name,
        "[Authorised Signatory]": business.get("legalName") or business.get("name") or "Safar Self Drive",
        "[LLPIN]": business.get("llpin") or "To be completed",
        "[Registered Office Address, Ahmedabad, Gujarat]": business.get("address") or "Ahmedabad, Gujarat, India",
        "[legal@safarcars.com]": business.get("email") or "support@safarcars.com",
        "[grievance@safarcars.com]": business.get("email") or "support@safarcars.com",
        "[date, time, IP]": f"{fmt_date(effective)} (Platform electronic log)",
        "[Booking ID]": str(booking.get("reference") or booking.get("id") or "Not recorded"),
        "[DD Month YYYY, hh:mm]": fmt_date(booking.get("startDate")),
        "[Phone / Email]": " / ".join(v for v in [customer.get("phone"), customer.get("email")] if v) or "Not recorded",
        "[DL number]": "Verified on Platform; number withheld",
        "[Yes / No]": "Yes" if customer.get("isVerified", True) else "No",
        "[Name, DL number, KYC status]": "None recorded",
        "[Make / Model / Variant / Fuel]": " / ".join(v for v in [vehicle_name, vehicle.get("fuelType")] if v),
        "[Reg. no.]": vehicle.get("registrationNumber") or "Not recorded",
        "[Chassis no.]": "Not recorded on listing",
        "[Engine no.]": "Not recorded on listing",
        "[Policy no. / Expiry]": "Verified vehicle insurance; policy details withheld",
        "[Address]": trip.get("pickupLocation") or vehicle.get("locationCity") or "Not recorded",
        "[₹ amount]": money(booking.get("totalAmount")),
        "[km per day / total km]": f"{booking.get('includedKilometres') or policy.get('includedKilometresPer24Hours') or 0} km per 24 hours",
        "[Same-to-same / Full-to-full]": "Same-to-same",
        "[Details, if any]": "As shown in the vehicle listing",
        "[21]": "21",
        "[1] year": "1 year",
        "[30] minutes": "30 minutes",
        "[3] days": "3 days",
        "[7] working days": "7 working days",
        "[30] days": "30 days",
        "[48]": "48",
        "[24]": "24",
        "[6]": "6",
        "[15] days": "15 days",
        "[3] years": "3 years",
        "[Safarcars / the Guest / the Host]": "Safarcars",
        "[Ahmedabad, Gujarat]": business.get("address") or "Ahmedabad, Gujarat",
    }

    document = Document(args.template)
    for paragraph in walk_paragraphs(document):
        replace_paragraph(paragraph, replacements)

    if len(document.tables) >= 1:
        signatures = document.tables[0]
        set_cell(signatures.cell(1, 0), f"Name: {host_name}")
        set_cell(signatures.cell(1, 1), f"Name: {customer_name}")
        set_cell(signatures.cell(1, 2), f"Name: {business.get('legalName') or business.get('name') or 'Safar Self Drive'}")
        accepted = f"Accepted on: {fmt_date(effective)} (Platform electronic log)"
        set_cell(signatures.cell(2, 0), accepted)
        set_cell(signatures.cell(2, 1), accepted)

    if len(document.tables) >= 2:
        schedule = document.tables[1]
        details = {
            "Booking ID": booking.get("reference") or booking.get("id"),
            "Booking Start Date and time": fmt_date(booking.get("startDate")),
            "Booking End Date and time": fmt_date(booking.get("endDate")),
            "Host name": host_name,
            "Primary Guest name": customer_name,
            "Primary Guest phone / email": " / ".join(v for v in [customer.get("phone"), customer.get("email")] if v),
            "Driving licence number (Primary Guest)": "Verified on Platform; number withheld",
            "KYC verified": "Yes" if customer.get("isVerified", True) else "No",
            "Co-Driver (if any)": "None recorded",
            "Vehicle make, model, variant, fuel": " / ".join(v for v in [vehicle_name, vehicle.get("fuelType")] if v),
            "Registration number": vehicle.get("registrationNumber") or "Not recorded",
            "Chassis number": "Not recorded on listing",
            "Engine number": "Not recorded on listing",
            "Insurance policy number and expiry": "Verified vehicle insurance; policy details withheld",
            "PUC valid until": "Verified through return date",
            "Designated Location (pick-up and drop-off)": trip.get("pickupLocation") or vehicle.get("locationCity") or "Not recorded",
            "Booking Fee": money(booking.get("rentalAmount")),
            "Security Deposit": money(booking.get("securityDeposit")),
            "Included Kilometres": f"{booking.get('includedKilometres') or 0} km per 24 hours",
            "Fuel policy": "Same-to-same",
            "Restricted routes or areas": "As shown in the vehicle listing",
        }
        for row in schedule.rows[1:]:
            label = cell_text(row.cells[0])
            if label in details:
                set_cell(row.cells[1], str(details[label]))

    if len(document.tables) >= 3:
        schedule = document.tables[2]
        charges = {
            "Platform convenience fee": "Included in booking checkout",
            "Trip protection / damage waiver (optional)": "Not selected",
            "Excess (own-damage, insurance-admitted claim)": "As per approved insurance claim",
            "Cap on non-payable items (depreciation, tyres, glass, battery)": "Actual verified loss",
            "Extra kilometre rate": money(policy.get("excessKmRate")) + " per km",
            "Late-return charge": money(policy.get("lateReturnRatePerHour")) + f" per hour after {policy.get('lateGraceMinutes', 30)} minutes",
            "Fuel shortfall": "Actual fuel cost plus applicable service charge",
            "Cleaning charge (smoking, spill, pet hair, strong odour)": "Actual documented cost",
            "Challan processing fee": "Actual documented cost",
            "Lost key / key fob": "Actual replacement cost",
            "Lost or damaged documents (duplicate and inconvenience fee)": "Actual documented cost",
            "Loss-of-use charge (repair, impound, detention)": "Actual documented cost",
            "Delivery or pick-up outside Designated Location": "As quoted at checkout",
            "Vehicle returned at a wrong location": "Transport cost plus applicable charge",
            "Replacement of In-Vehicle Device (tampering)": "Actual replacement cost",
            "Late payment charge on dues": "As per published Fee Policy",
        }
        for row in schedule.rows[1:]:
            label = cell_text(row.cells[0])
            if label in charges:
                set_cell(row.cells[1], charges[label])

    if len(document.tables) >= 4:
        schedule = document.tables[3]
        refunds = {
            "More than 48 hours before start": "100% (less Platform fee, if any)",
            "Between 24 and 48 hours before start": "50%",
            "Between 6 and 24 hours before start": "0%",
            "Less than 6 hours before start, or no-show": "0% (no refund)",
            "Cancelled by Platform for failed KYC or verification": "As per Clause 3.3",
            "Vehicle not delivered or unsafe at delivery": "100% refund and Platform assistance",
        }
        for row in schedule.rows[1:]:
            label = cell_text(row.cells[0])
            if label in refunds:
                set_cell(row.cells[1], refunds[label])

    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)


if __name__ == "__main__":
    main()
