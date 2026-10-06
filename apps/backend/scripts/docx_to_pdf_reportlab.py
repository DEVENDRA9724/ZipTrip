"""Portable PDF fallback for the supplied agreement template.

It keeps the source document's wording, section order, schedules and populated
values when a native Word/LibreOffice converter is unavailable.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def register_fonts() -> tuple[str, str]:
    regular = "Helvetica"
    bold = "Helvetica-Bold"
    for path, name, is_bold in [
        (r"C:\Windows\Fonts\arial.ttf", "SafarArial", False),
        (r"C:\Windows\Fonts\arialbd.ttf", "SafarArialBold", True),
    ]:
        if Path(path).exists():
            pdfmetrics.registerFont(TTFont(name, path))
            if is_bold:
                bold = name
            else:
                regular = name
    return regular, bold


def clean(text: str) -> str:
    return escape(" ".join(text.split())).replace("\n", "<br/>")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    document = Document(args.input)
    regular, bold = register_fonts()
    styles = getSampleStyleSheet()
    body = ParagraphStyle("SafarBody", parent=styles["BodyText"], fontName=regular, fontSize=9.2, leading=13, spaceAfter=5)
    title = ParagraphStyle("SafarTitle", parent=body, fontName=bold, fontSize=17, leading=21, alignment=TA_CENTER, spaceAfter=8)
    heading = ParagraphStyle("SafarHeading", parent=body, fontName=bold, fontSize=11.5, leading=15, spaceBefore=7, spaceAfter=4)
    small = ParagraphStyle("SafarSmall", parent=body, fontSize=8.2, leading=10.5)

    paragraphs = {id(p._p): p for p in document.paragraphs}
    tables = {id(t._tbl): t for t in document.tables}
    story = []
    first_text = True
    for node in document.element.body.iterchildren():
        if node.tag.endswith("}p") and id(node) in paragraphs:
            text = paragraphs[id(node)].text.strip()
            if not text:
                story.append(Spacer(1, 2 * mm))
                continue
            if first_text:
                story.append(Paragraph(clean(text), title))
                first_text = False
            elif text.isupper() or text.startswith("SCHEDULE") or re.match(r"^(\d+\.|WHEREAS|NOW,|IN WITNESS)", text):
                story.append(Paragraph(clean(text), heading))
            else:
                story.append(Paragraph(clean(text), body))
        elif node.tag.endswith("}tbl") and id(node) in tables:
            source = tables[id(node)]
            rows = []
            for row in source.rows:
                values = [Paragraph(clean(cell.text), small) for cell in row.cells]
                rows.append(values)
            if rows:
                count = len(rows[0])
                widths = [52 * mm, 92 * mm, 35 * mm] if count == 3 else [58 * mm, 121 * mm]
                table = Table(rows, colWidths=widths[:count], repeatRows=1 if len(rows) > 1 else 0, hAlign="LEFT")
                table.setStyle(TableStyle([
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D5D9DF")),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF1F5")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]))
                story.extend([Spacer(1, 2 * mm), table, Spacer(1, 3 * mm)])

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    pdf = SimpleDocTemplate(str(output), pagesize=A4, rightMargin=17 * mm, leftMargin=17 * mm, topMargin=16 * mm, bottomMargin=16 * mm, title="Safarcars Car Sharing Agreement")
    pdf.build(story)


if __name__ == "__main__":
    main()
